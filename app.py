"""SP 广告周报服务：POST /weekly → 读飞书表 → 生成 Excel → 上传飞书 → 归档表 + 群卡片
环境变量：FS_APP_ID, FS_APP_SECRET, FS_WEBHOOK(群机器人), FS_FOLDER(云文档文件夹token), API_KEY(调用口令, 可选), PORT
"""
import os, sys, json, datetime, tempfile, subprocess, shutil, threading, traceback, time
from flask import Flask, request, jsonify
import feishu as fs
from data import Data, APP, T
import report, suggest

app = Flask(__name__)
LOCK = threading.Lock()
LAST = {'status': 'idle'}
WEBHOOK = os.environ.get('FS_WEBHOOK', '')
FOLDER = os.environ.get('FS_FOLDER', '')
API_KEY = os.environ.get('API_KEY', '')
OUT_DIR = os.environ.get('OUT_DIR', '/tmp/reports')
BASE_URL = f'https://kcn46c3x2v7r.feishu.cn/base/{APP}'


def periods(today=None, weeks=5):
    """成熟期 = 最近 5 个完整周（截至 today-8 天所在周之前的周日）；新鲜周 = 紧接着的一周"""
    today = today or (datetime.datetime.utcnow() + datetime.timedelta(hours=8)).date()
    last_sun = today - datetime.timedelta(days=today.weekday() + 1)  # 上一个周日
    end = last_sun
    mature_end = end - datetime.timedelta(days=7)
    start = mature_end - datetime.timedelta(days=7 * weeks - 1)
    WEEKS = []
    a = start
    while a <= end:
        WEEKS.append((a, a, a + datetime.timedelta(days=6))); a += datetime.timedelta(days=7)
    return {'snap': today, 'start': start, 'end': end, 'mature_end': mature_end, 'weeks': WEEKS, 'mature_weeks': weeks, 'week_of': end - datetime.timedelta(days=6)}


def recalc(path, timeout=240):
    """LibreOffice 无头重算公式（写入缓存值，飞书/手机预览才有数字）"""
    if not shutil.which('soffice'): return 'soffice missing'
    prof = tempfile.mkdtemp(prefix='lo-')
    try:
        subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{prof}', '--terminate_after_init'], capture_output=True, timeout=60)
        mdir = os.path.join(prof, 'user', 'basic', 'Standard'); os.makedirs(mdir, exist_ok=True)
        open(os.path.join(mdir, 'Module1.xba'), 'w').write('<?xml version="1.0" encoding="UTF-8"?>\n<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">\n<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">\nSub RecalculateAndSave()\n  ThisComponent.calculateAll()\n  ThisComponent.store()\n  ThisComponent.close(True)\nEnd Sub\n</script:module>')
        lib = os.path.join(mdir, 'script.xlb')
        if not os.path.exists(lib): open(lib, 'w').write('<?xml version="1.0" encoding="UTF-8"?>\n<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">\n<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false">\n <library:element library:name="Module1"/>\n</library:library>')
        r = subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{prof}', 'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application', os.path.abspath(path)], capture_output=True, text=True, timeout=timeout)
        return 'ok' if r.returncode == 0 else f'rc={r.returncode} {r.stderr[-300:]}'
    except Exception as e:
        return f'recalc error: {e}'
    finally:
        shutil.rmtree(prof, ignore_errors=True)


def check_errors(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True); n = 0; f = 0; ex = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?') in ('#VALUE', '#DIV/0', '#REF', '#NAME', '#N/A', '#NUM', '#NULL'):
                    n += 1; ex.append(f'{ws.title}!{c.coordinate}') if len(ex) < 10 else None
    return n, ex


def summarize(d, P):
    """卡片摘要：两店成熟期 vs 上一个5周？简化为成熟期合计 + 新鲜周花费/CPC"""
    out = []
    for st in ('HLZ-US', 'YYK-US'):
        m = [r for r in d.camp if r['店铺'] == st and P['start'] <= r['报表日期'] <= P['mature_end']]
        f = [r for r in d.camp if r['店铺'] == st and r['报表日期'] > P['mature_end']]
        if not m and not f: continue
        mc = sum(r['花费'] for r in m); ms = sum(r['销售额'] for r in m); mk = sum(r['点击量'] for r in m)
        fc = sum(r['花费'] for r in f); fk = sum(r['点击量'] for r in f); f1 = sum(r.get('销售额1d', 0) for r in f); m1 = sum(r.get('销售额1d', 0) for r in m)
        acos = mc / ms if ms else 0; cpc = mc / mk if mk else 0; fcpc = fc / fk if fk else 0
        a1 = fc / f1 if f1 else 0; am1 = mc / m1 if m1 else 0
        out.append(f"**{st}** 成熟期5周：花费 ${mc:,.0f}｜销售 ${ms:,.0f}｜ACOS {acos * 100:.1f}%｜CPC ${cpc:.2f}\n　新鲜周：花费 ${fc:,.0f}（周均 ${mc / 5:,.0f}）｜CPC ${fcpc:.2f}｜1d口径ACOS {a1 * 100:.0f}%（成熟期同口径 {am1 * 100:.0f}%）")
    return out


def run(dry=False, today=None, no_send=False, cache_dir=None):
    P = periods(today); P.update({'default_target': float(os.environ.get('DEFAULT_TARGET_ACOS', 0.30)), 'zero_cost': 20, 'min_clicks': 100})
    t0 = time.time()
    d = Data(P['start'], P['end']).load(cache_dir)
    sug = suggest.generate(d, P)
    n_new, n_upd, rows = suggest.sync(d, P, sug, dry=dry)
    P['opt_rows'] = rows
    os.makedirs(OUT_DIR, exist_ok=True)
    fname = f"SP广告周报_{P['snap'].strftime('%Y%m%d')}_成熟{P['start'].strftime('%m%d')}-{P['mature_end'].strftime('%m%d')}.xlsx"
    path = os.path.join(OUT_DIR, fname)
    stats = report.build(d, P, path)
    rc = recalc(path); n_err, ex = check_errors(path)
    res = {'file': fname, 'path': path, 'periods': {k: str(v) for k, v in P.items() if k in ('snap', 'start', 'end', 'mature_end', 'week_of')}, 'stats': stats, 'suggest_new': n_new, 'suggest_refreshed': n_upd, 'recalc': rc, 'formula_errors': n_err, 'error_cells': ex, 'seconds': round(time.time() - t0)}
    if dry or no_send: return res
    tok = fs.upload_file(path, FOLDER, fname); link = fs.share_link(tok)
    pending = sum(1 for o in rows if o.get('状态') == '待处理')
    fs.batch_create(APP, T['arch'], [{'fields': {'文件名': fname, '生成日期': fs.date_ms(P['snap']), '成熟期': f"{P['start']}~{P['mature_end']}", '新鲜周': f"{P['week_of']}~{P['end']}", '链接': {'link': link, 'text': fname}, '摘要': '\n'.join(summarize(d, P)).replace('**', ''), '新增建议数': n_new, '待处理建议数': pending, '文件token': tok}}])
    if WEBHOOK:
        lines = summarize(d, P) + [f"本周新增建议 {n_new} 条，待处理共 {pending} 条 → [优化记录]({BASE_URL}?table={T['opt']})", f"[历史周报归档]({BASE_URL}?table={T['arch']})"]
        fs.send_card(WEBHOOK, f"SP广告周报 · {P['snap']}（成熟期 {P['start'].strftime('%m/%d')}–{P['mature_end'].strftime('%m/%d')}）", lines, link, '打开本周 Excel')
    res.update({'link': link, 'token': tok, 'pending': pending})
    return res


def _bg(kw):
    with LOCK:
        LAST.update({'status': 'running', 'started': str(datetime.datetime.utcnow())})
        try:
            LAST.update({'status': 'done', 'result': run(**kw)})
        except Exception as e:
            LAST.update({'status': 'error', 'error': str(e), 'trace': traceback.format_exc()[-2000:]})
            if WEBHOOK and not kw.get('dry'):
                try: fs.send_card(WEBHOOK, 'SP广告周报生成失败', [f'错误：{e}'], color='red')
                except Exception: pass


@app.route('/health')
def health(): return jsonify({'ok': True, 'last': LAST})


@app.route('/weekly', methods=['POST', 'GET'])
def weekly():
    if API_KEY and request.headers.get('X-Api-Key', request.args.get('key')) != API_KEY: return jsonify({'error': 'unauthorized'}), 401
    q = request.get_json(silent=True) or {}; q.update(request.args.to_dict())
    kw = {'dry': str(q.get('dryRun', '')).lower() in ('1', 'true'), 'no_send': str(q.get('noSend', '')).lower() in ('1', 'true')}
    if q.get('today'): kw['today'] = datetime.date.fromisoformat(q['today'])
    if str(q.get('wait', '')).lower() in ('1', 'true'):
        try: return jsonify(run(**kw))
        except Exception as e: return jsonify({'error': str(e), 'trace': traceback.format_exc()[-2000:]}), 500
    if LOCK.locked(): return jsonify({'status': 'already running', 'last': LAST}), 409
    threading.Thread(target=_bg, args=(kw,), daemon=True).start()
    return jsonify({'status': 'started', 'params': kw})


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'run':
        kw = {'dry': '--dry' in sys.argv, 'no_send': '--no-send' in sys.argv}
        for a in sys.argv:
            if a.startswith('--today='): kw['today'] = datetime.date.fromisoformat(a.split('=')[1])
            if a.startswith('--cache='): kw['cache_dir'] = a.split('=')[1]
        print(json.dumps(run(**kw), ensure_ascii=False, indent=1))
    else:
        app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8080)))
